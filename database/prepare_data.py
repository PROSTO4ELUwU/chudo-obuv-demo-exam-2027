"""Подготовка данных заказчика к импорту и сборка итогового скрипта БД.

Скрипт читает файлы *_import.xlsx, исправляет найденные в них ошибки,
проверяет ссылочную целостность и формирует chudo_obuv.sql:
структура из schema.sql и данные в виде INSERT-запросов.

Запуск из корня репозитория:
    python database/prepare_data.py
"""

from collections.abc import Iterable
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

DATABASE_DIR = Path(__file__).resolve().parent
IMPORT_DIR = DATABASE_DIR / "import"
IMAGES_DIR = DATABASE_DIR.parent / "resources" / "images"
SCHEMA_FILE = DATABASE_DIR / "schema.sql"
OUTPUT_FILE = DATABASE_DIR / "chudo_obuv.sql"

ENCODING_LINE = "SET client_encoding = 'UTF8';"

# В остатках и заказах эта модель записана короче, чем в каталоге.
# Эталонным источником наименований считается каталог (Products_import).
PRODUCT_NAME_FIXES = {
    ("Черные туфли в классическом стиле", "Барбари"):
        "Черные туфли в классическом стиле — база для деловых образов",
}

ProductKey = tuple[str, str]


class DataError(Exception):
    """Ошибка в данных заказчика, которую нельзя исправить автоматически."""


@dataclass(frozen=True)
class User:
    """Пользователь системы."""

    last_name: str
    first_name: str
    patronymic: str | None
    login: str
    role_name: str

    @property
    def full_name(self) -> str:
        """Фамилия, имя и отчество через пробел."""
        parts = (self.last_name, self.first_name, self.patronymic)
        return " ".join(part for part in parts if part)


@dataclass(frozen=True)
class Product:
    """Модель обуви из каталога."""

    name: str
    manufacturer: str
    category: str
    subcategory: str
    description: str
    composition: str
    price: Decimal
    image_file: str | None

    @property
    def key(self) -> ProductKey:
        """Естественный ключ модели: наименование и производство."""
        return self.name, self.manufacturer


@dataclass(frozen=True)
class StockItem:
    """Товарная позиция: модель, размер и доступное количество."""

    product_key: ProductKey
    size: Decimal
    quantity: int


@dataclass(frozen=True)
class OrderItem:
    """Строка заказа."""

    product_key: ProductKey
    size: Decimal
    quantity: int
    unit_price: Decimal


@dataclass
class Order:
    """Заказ с номером из файла заказчика."""

    number: int
    order_date: date
    client_login: str
    items: list[OrderItem] = field(default_factory=list)


def clean_text(value: object) -> str:
    """Возвращает текст без лишних пробелов.

    split() без аргументов делит строку и по неразрывному пробелу (U+00A0),
    который попал в наименования из Excel.
    """
    if value is None:
        return ""
    return " ".join(str(value).split())


def to_size(value: object) -> Decimal:
    """Размер обуви с точностью до половины: 36.5."""
    return Decimal(str(value)).quantize(Decimal("0.1"))


def to_money(value: object) -> Decimal:
    """Сумма в рублях с копейками."""
    return Decimal(str(value)).quantize(Decimal("0.01"))


def to_date(value: object) -> date:
    """Дата из ячейки Excel."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raise DataError(f"Ожидалась дата, получено «{value}»")


def ensure_unique(values: list, what: str) -> None:
    """Проверяет, что в списке нет повторов, иначе сообщает о них."""
    duplicates = sorted({str(value) for value in values if values.count(value) > 1})
    if duplicates:
        raise DataError(f"Повторяется {what}: {', '.join(duplicates)}")


def unique_in_order(values: Iterable[str]) -> list[str]:
    """Уникальные значения в порядке первого появления."""
    return list(dict.fromkeys(values))


def read_sheet(file_name: str) -> list[dict[str, object]]:
    """Читает первый лист файла в список словарей «заголовок → значение»."""
    workbook = load_workbook(IMPORT_DIR / file_name, data_only=True)
    rows = workbook.active.iter_rows(values_only=True)
    header = [clean_text(title) for title in next(rows)]
    return [dict(zip(header, row)) for row in rows if any(cell is not None for cell in row)]


def load_users() -> list[User]:
    """Загружает пользователей и проверяет уникальность логинов и ФИО."""
    users = [
        User(
            last_name=clean_text(row["Фамилия"]),
            first_name=clean_text(row["Имя"]),
            # Опечатка в заголовке столбца сохранена как в файле заказчика
            patronymic=clean_text(row["Отчетсво"]) or None,
            login=clean_text(row["Логин"]),
            role_name=clean_text(row["Роль"]),
        )
        for row in read_sheet("Users_import.xlsx")
    ]
    ensure_unique([user.login for user in users], "логин")
    ensure_unique([user.full_name for user in users], "ФИО пользователя")
    return users


def load_products(fixes: set[str]) -> list[Product]:
    """Загружает каталог; текст всех полей очищается от лишних пробелов."""
    products = []
    for row in read_sheet("Products_import.xlsx"):
        raw_name = str(row["Наименование товара"])
        name = clean_text(raw_name)
        if name != raw_name:
            fixes.add(f"Products_import: лишние или неразрывные пробелы в «{name}»")
        image_file = clean_text(row["Изображение"]) or None
        if image_file and not (IMAGES_DIR / image_file).is_file():
            fixes.add(f"Products_import: нет файла {image_file}, будет показана заглушка")
            image_file = None
        products.append(Product(
            name=name,
            manufacturer=clean_text(row["Производство"]),
            category=clean_text(row["Категория"]),
            subcategory=clean_text(row["Подкатегория"]),
            description=clean_text(row["Описание"]),
            composition=clean_text(row["Состав"]),
            price=to_money(row["Цена"]),
            image_file=image_file,
        ))
    ensure_unique([product.key for product in products], "модель в каталоге")
    return products


def load_sizes() -> list[Decimal]:
    """Загружает размерную сетку."""
    sizes = [to_size(row["Размер"]) for row in read_sheet("Sizes_import.xlsx")]
    ensure_unique(sizes, "размер")
    return sizes


def resolve_product(row: dict[str, object], catalog: dict[ProductKey, Product],
                    source: str, fixes: set[str]) -> Product:
    """Находит модель каталога, на которую ссылается строка остатков или заказов."""
    key = (clean_text(row["Наименование товара"]), clean_text(row["Производство"]))
    if key in PRODUCT_NAME_FIXES:
        fixed_key = (PRODUCT_NAME_FIXES[key], key[1])
        fixes.add(f"{source}: «{key[0]}» сопоставлено с моделью каталога «{fixed_key[0]}»")
        key = fixed_key
    if key not in catalog:
        raise DataError(f"{source}: модель «{key[0]}» ({key[1]}) не найдена в каталоге")
    return catalog[key]


def load_stock_items(catalog: dict[ProductKey, Product], sizes: set[Decimal],
                     fixes: set[str]) -> list[StockItem]:
    """Загружает товарные позиции и проверяет ссылки на каталог и размеры."""
    items = []
    for row in read_sheet("Stock_Items_import.xlsx"):
        product = resolve_product(row, catalog, "Stock_Items_import", fixes)
        size = to_size(row["Размер"])
        if size not in sizes:
            raise DataError(f"Stock_Items_import: размера {size} нет в размерной сетке")
        items.append(StockItem(product.key, size, int(row["Количество доступное для заказа"])))
    ensure_unique([(item.product_key, item.size) for item in items], "товарная позиция")
    return items


def load_orders(catalog: dict[ProductKey, Product], stock_keys: set[tuple[ProductKey, Decimal]],
                users: list[User], fixes: set[str]) -> list[Order]:
    """Собирает заказы из строк файла: одна строка — одна позиция заказа."""
    logins = {user.full_name: user.login for user in users}
    orders: dict[int, Order] = {}
    for row in read_sheet("Orders_import.xlsx"):
        number = int(row["Номер заказа"])
        client_name = clean_text(row["ФИО"])
        if client_name not in logins:
            raise DataError(f"Заказ №{number}: клиент «{client_name}» не найден среди пользователей")
        order_date = to_date(row["Дата заказа"])
        order = orders.setdefault(number, Order(number, order_date, logins[client_name]))
        if (order.order_date, order.client_login) != (order_date, logins[client_name]):
            raise DataError(f"Заказ №{number}: строки заказа расходятся по дате или клиенту")

        product = resolve_product(row, catalog, "Orders_import", fixes)
        if product.category != clean_text(row["Категория"]):
            raise DataError(f"Заказ №{number}: категория «{product.name}» не совпадает с каталогом")
        size = to_size(row["Размер"])
        if (product.key, size) not in stock_keys:
            raise DataError(f"Заказ №{number}: нет товарной позиции «{product.name}», размер {size}")
        order.items.append(OrderItem(product.key, size, int(row["Количество"]),
                                     to_money(row["Цена за единицу"])))

    for order in orders.values():
        positions = [(item.product_key, item.size) for item in order.items]
        ensure_unique(positions, f"позиция в заказе №{order.number}")
    return sorted(orders.values(), key=lambda order: order.number)


def sql_literal(value: object) -> str:
    """Записывает значение Python как литерал SQL."""
    if value is None:
        return "NULL"
    if isinstance(value, (int, Decimal)):
        return str(value)
    if isinstance(value, date):
        return f"DATE '{value.isoformat()}'"
    return "'" + str(value).replace("'", "''") + "'"


def values_list(rows: Iterable[tuple]) -> str:
    """Строки конструкции VALUES — по одной записи на строку."""
    return ",\n".join(
        "    (" + ", ".join(sql_literal(value) for value in row) + ")" for row in rows
    )


def numbered(rows: Iterable[tuple]) -> list[tuple]:
    """Добавляет в начало каждой записи её номер по порядку в файле заказчика."""
    return [(number, *row) for number, row in enumerate(rows, start=1)]


def build_data_sql(users: list[User], products: list[Product], sizes: list[Decimal],
                   stock_items: list[StockItem], orders: list[Order]) -> str:
    """Формирует INSERT-запросы; внешние ключи находятся по естественным ключам."""
    roles = unique_in_order(user.role_name for user in users)
    categories = unique_in_order(product.category for product in products)
    subcategories = unique_in_order(product.subcategory for product in products)
    manufacturers = unique_in_order(product.manufacturer for product in products)
    order_items = [
        (order.number, *item.product_key, item.size, item.quantity, item.unit_price)
        for order in orders
        for item in order.items
    ]
    return f"""\
-- ---------------------------------------------------------------------------
-- Данные
-- Столбец n — номер строки в файле заказчика: соединения (JOIN) меняют
-- порядок строк, а ORDER BY v.n сохраняет его в идентификаторах.
-- ---------------------------------------------------------------------------

INSERT INTO roles (role_name) VALUES
{values_list((role,) for role in roles)};

INSERT INTO categories (category_name) VALUES
{values_list((category,) for category in categories)};

INSERT INTO subcategories (subcategory_name) VALUES
{values_list((subcategory,) for subcategory in subcategories)};

INSERT INTO manufacturers (manufacturer_name) VALUES
{values_list((manufacturer,) for manufacturer in manufacturers)};

INSERT INTO sizes (size_value) VALUES
{values_list((size,) for size in sizes)};

INSERT INTO users (last_name, first_name, patronymic, login, role_id)
SELECT v.last_name, v.first_name, v.patronymic, v.login, r.role_id
FROM (VALUES
{values_list(numbered((u.last_name, u.first_name, u.patronymic, u.login, u.role_name)
                      for u in users))}
) AS v (n, last_name, first_name, patronymic, login, role_name)
JOIN roles AS r ON r.role_name = v.role_name
ORDER BY v.n;

INSERT INTO products (product_name, category_id, subcategory_id, manufacturer_id,
                      description, composition, price, image_file)
SELECT v.product_name, c.category_id, s.subcategory_id, m.manufacturer_id,
       v.description, v.composition, v.price, v.image_file
FROM (VALUES
{values_list(numbered((p.name, p.category, p.subcategory, p.manufacturer, p.description,
                       p.composition, p.price, p.image_file) for p in products))}
) AS v (n, product_name, category_name, subcategory_name, manufacturer_name,
        description, composition, price, image_file)
JOIN categories AS c ON c.category_name = v.category_name
JOIN subcategories AS s ON s.subcategory_name = v.subcategory_name
JOIN manufacturers AS m ON m.manufacturer_name = v.manufacturer_name
ORDER BY v.n;

INSERT INTO stock_items (product_id, size_id, quantity)
SELECT p.product_id, sz.size_id, v.quantity
FROM (VALUES
{values_list(numbered((*item.product_key, item.size, item.quantity) for item in stock_items))}
) AS v (n, product_name, manufacturer_name, size_value, quantity)
JOIN manufacturers AS m ON m.manufacturer_name = v.manufacturer_name
JOIN products AS p ON p.product_name = v.product_name AND p.manufacturer_id = m.manufacturer_id
JOIN sizes AS sz ON sz.size_value = v.size_value
ORDER BY v.n;

-- Номера заказов сохраняются такими же, как в файле заказчика
INSERT INTO orders (order_id, order_date, client_id)
OVERRIDING SYSTEM VALUE
SELECT v.order_id, v.order_date, u.user_id
FROM (VALUES
{values_list((order.number, order.order_date, order.client_login) for order in orders)}
) AS v (order_id, order_date, login)
JOIN users AS u ON u.login = v.login
ORDER BY v.order_id;

-- Новые заказы продолжат нумерацию после последнего перенесённого
SELECT setval(pg_get_serial_sequence('orders', 'order_id'), (SELECT max(order_id) FROM orders));

INSERT INTO order_items (order_id, stock_item_id, quantity, unit_price)
SELECT v.order_id, si.stock_item_id, v.quantity, v.unit_price
FROM (VALUES
{values_list(numbered(order_items))}
) AS v (n, order_id, product_name, manufacturer_name, size_value, quantity, unit_price)
JOIN manufacturers AS m ON m.manufacturer_name = v.manufacturer_name
JOIN products AS p ON p.product_name = v.product_name AND p.manufacturer_id = m.manufacturer_id
JOIN sizes AS sz ON sz.size_value = v.size_value
JOIN stock_items AS si ON si.product_id = p.product_id AND si.size_id = sz.size_id
ORDER BY v.n;

-- Строка с неверным наименованием молча отбросилась бы соединением (JOIN),
-- поэтому количество загруженных строк сверяется с файлами заказчика
DO $$
BEGIN
    IF (SELECT count(*) FROM users) <> {len(users)}
        OR (SELECT count(*) FROM products) <> {len(products)}
        OR (SELECT count(*) FROM stock_items) <> {len(stock_items)}
        OR (SELECT count(*) FROM orders) <> {len(orders)}
        OR (SELECT count(*) FROM order_items) <> {len(order_items)}
    THEN
        RAISE EXCEPTION 'Загружены не все строки данных заказчика';
    END IF;
END $$;
"""


def build_script(data_sql: str) -> str:
    """Собирает итоговый скрипт: структура и данные в одной транзакции."""
    schema_lines = SCHEMA_FILE.read_text(encoding="utf-8").splitlines()
    if not schema_lines[0].startswith(ENCODING_LINE):
        raise DataError(f"schema.sql должен начинаться со строки {ENCODING_LINE}")
    schema_body = "\n".join(schema_lines[1:]).strip()
    return f"""\
{ENCODING_LINE}  -- psql в русской Windows иначе читает файл как WIN1251

-- ============================================================================
-- Скрипт базы данных «Чудо Обувь»: структура и данные
-- Сформирован автоматически скриптом database/prepare_data.py
-- из файлов заказчика database/import/*_import.xlsx
-- ============================================================================

BEGIN;

{schema_body}

{data_sql}
COMMIT;
"""


def main() -> None:
    """Готовит данные и записывает chudo_obuv.sql."""
    fixes: set[str] = set()
    try:
        users = load_users()
        products = load_products(fixes)
        sizes = load_sizes()
        catalog = {product.key: product for product in products}
        stock_items = load_stock_items(catalog, set(sizes), fixes)
        stock_keys = {(item.product_key, item.size) for item in stock_items}
        orders = load_orders(catalog, stock_keys, users, fixes)
        script = build_script(build_data_sql(users, products, sizes, stock_items, orders))
    except DataError as error:
        raise SystemExit(f"Ошибка в данных: {error}") from error

    OUTPUT_FILE.write_text(script, encoding="utf-8", newline="\n")
    order_items_count = sum(len(order.items) for order in orders)
    print(f"Пользователей: {len(users)}, моделей: {len(products)}, размеров: {len(sizes)}, "
          f"товарных позиций: {len(stock_items)}, заказов: {len(orders)} "
          f"({order_items_count} позиций)")
    print("Исправления в данных заказчика:")
    for fix in sorted(fixes):
        print(f"  - {fix}")
    print(f"Скрипт записан: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
